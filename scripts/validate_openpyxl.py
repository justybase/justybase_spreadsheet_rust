#!/usr/bin/env python3
"""Independent XLSX validation with openpyxl (a third-party, pure-Python reader).

This is the openpyxl leg of the write/read correctness process: it proves that
files produced by this crate's XlsxWriter are not just readable by our own
XlsxReader and by calamine, but also by the library Excel tooling around
Python (openpyxl is what most Python pipelines use).

What it checks
--------------
1. Value-level parity on the `parity_write` fixtures (types, exact values,
   number formats, sheet layout, hidden sheet flag, empty sheets):
     - strings survive XML escaping (Polish diacritics, &, <>, quotes)
     - integers/floats/booleans keep their Excel types
     - datetimes round-trip (<= 1 s tolerance for serial rounding)
     - formatted cells carry the right number formats
     - > 2^53 integers stay TEXT (no precision loss)
     - NaN stays the string "NaN"; empty string cells are kept
2. Optional dimension checks for larger outputs, e.g. the write benchmarks:

     python3 scripts/validate_openpyxl.py \
         /tmp/parity_check/spreadsheet-batch.xlsx=5001x7 \
         /tmp/parity_check/spreadsheet-streaming.xlsx=5001x7

Notes
-----
- openpyxl reads .xlsx only; .xlsb coverage comes from the XlsbReader
  roundtrip tests and the calamine cross-check in `read_comparison`.
- If openpyxl is missing, the script bootstraps a virtualenv under
  `target/ovenv` (gitignored) and re-executes itself with it.

Exit status: 0 on success, 1 on any validation failure.
"""

from __future__ import annotations

import datetime as dt
import os
import subprocess
import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
VENV_DIR = PROJECT_ROOT / "target" / "ovenv"

DATETIME_TOLERANCE_S = 1.0
FLOAT_TOLERANCE = 1e-9


def bootstrap_openpyxl() -> None:
    """Re-execute under a project venv when openpyxl is not importable."""
    try:
        import openpyxl  # noqa: F401

        return
    except ImportError:
        pass

    venv_python = VENV_DIR / ("Scripts" if os.name == "nt" else "bin") / (
        "python.exe" if os.name == "nt" else "python"
    )
    if not venv_python.exists():
        print(f"bootstrapping venv at {VENV_DIR} ...")
        subprocess.run([sys.executable, "-m", "venv", str(VENV_DIR)], check=True)
    if subprocess.run(
        [str(venv_python), "-c", "import openpyxl"], capture_output=True
    ).returncode != 0:
        print("installing openpyxl ...")
        subprocess.run(
            [str(venv_python), "-m", "pip", "install", "-q", "openpyxl>=3.1"],
            check=True,
        )
    os.execv(str(venv_python), [str(venv_python)] + sys.argv)


def generate_fixtures(out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    cmd = [
        "cargo",
        "run",
        "--quiet",
        "--release",
        "--example",
        "parity_write",
        str(out_dir),
    ]
    result = subprocess.run(cmd, cwd=PROJECT_ROOT, capture_output=True, text=True)
    if result.returncode != 0:
        sys.stderr.write(result.stdout + result.stderr)
        raise SystemExit("parity_write failed")


class Checker:
    def __init__(self) -> None:
        self.checks = 0
        self.failures: list[str] = []

    def expect(self, label: str, actual, expected) -> None:
        self.checks += 1
        # Floats go through decimal formatting (ryu) so compare with a
        # tolerance; e.g. -0.001 is not exactly representable in binary.
        if isinstance(actual, float) and isinstance(expected, float):
            if abs(actual - expected) > FLOAT_TOLERANCE:
                self.failures.append(f"{label}: expected {expected!r}, got {actual!r}")
            return
        if actual != expected:
            self.failures.append(f"{label}: expected {expected!r}, got {actual!r}")

    def expect_datetime(self, label: str, actual, expected: dt.datetime) -> None:
        """Excel serial rounding can shift sub-second; bound the error."""
        self.checks += 1
        if not isinstance(actual, dt.datetime):
            self.failures.append(f"{label}: expected datetime, got {actual!r}")
            return
        if abs((actual - expected).total_seconds()) > DATETIME_TOLERANCE_S:
            self.failures.append(f"{label}: expected {expected!r}, got {actual!r}")

    def expect_cell(
        self,
        ws,
        coord: str,
        value,
        number_format: str | None = None,
        kind: type | None = None,
    ) -> None:
        cell = ws[coord]
        self.expect(f"{ws.title}!{coord} value", cell.value, value)
        if kind is not None:
            self.expect(f"{ws.title}!{coord} type", type(cell.value), kind)
        if number_format is not None:
            self.expect(
                f"{ws.title}!{coord} number_format", cell.number_format, number_format
            )


def validate_basic(wb, ck: Checker) -> None:
    ck.expect("basic sheetnames", wb.sheetnames, ["Sheet1"])
    ws = wb["Sheet1"]
    ck.expect("basic dimensions", ws.dimensions, "A1:I4")

    headers = ["ID", "Name", "Count", "Score", "Active", "Joined", "Amount", "Big", "Missing"]
    for col, header in enumerate(headers, start=1):
        ck.expect_cell(ws, f"{ws.cell(1, col).coordinate}", header, kind=str)

    # Row 2 --------------------------------------------------------------
    ck.expect_cell(ws, "A2", 1, kind=int)
    ck.expect_cell(ws, "B2", "Alicja żółć", kind=str)  # diacritics via SST
    ck.expect_cell(ws, "C2", 42, kind=int)
    ck.expect_cell(ws, "D2", 12.5, kind=float)
    ck.expect_cell(ws, "E2", True, kind=bool)
    ck.expect_datetime("F2 datetime", ws["F2"].value, dt.datetime(2024, 1, 15, 12, 30, 45))
    ck.expect_cell(ws, "G2", 1234.5, number_format="#,##0.00", kind=float)
    # > 2^53 must stay TEXT so no precision is lost in Excel either.
    ck.expect_cell(ws, "H2", "9007199254740993", kind=str)
    ck.expect_cell(ws, "I2", None)

    # Row 3 --------------------------------------------------------------
    ck.expect_cell(ws, "A3", -7, kind=int)
    ck.expect_cell(ws, "B3", 'Bob & <Co> "q" \'x\'', kind=str)  # XML escaping
    ck.expect_cell(ws, "C3", 1 << 29, kind=int)
    ck.expect_cell(ws, "D3", -0.001, kind=float)
    ck.expect_cell(ws, "E3", False, kind=bool)
    ck.expect_cell(ws, "F3", dt.datetime(2023, 12, 31, 8, 0), number_format="yyyy-mm-dd")
    ck.expect_cell(ws, "G3", "  spaced  ", kind=str)  # xml:space="preserve"
    ck.expect_cell(ws, "H3", "0", kind=str)  # text "0", not number 0
    ck.expect_cell(ws, "I3", None)

    # Row 4 --------------------------------------------------------------
    ck.expect_cell(ws, "A4", 0, kind=int)
    ck.expect_cell(ws, "B4", "", kind=str)  # empty string cell survives
    ck.expect_cell(ws, "C4", -(1 << 29), kind=int)
    ck.expect_cell(ws, "D4", "NaN", kind=str)  # NaN stored as text
    ck.expect_cell(ws, "E4", True, kind=bool)
    # Serial 0 with a date style: openpyxl surfaces it as time(0,0) (or the
    #1899-12-30 epoch); accept either representation.
    ck.checks += 1
    if ws["F4"].value not in (dt.time(0, 0), dt.datetime(1899, 12, 30, 0, 0)):
        ck.failures.append(
            f"Sheet1!F4: expected midnight epoch, got {ws['F4'].value!r}"
        )
    ck.expect_cell(ws, "G4", 5, number_format='#,##0.00 "zł"', kind=int)
    ck.expect_cell(ws, "H4", "-123456789012345678901234567890", kind=str)
    ck.expect_cell(ws, "I4", None)


def validate_multi(wb, ck: Checker) -> None:
    ck.expect("multi sheetnames", wb.sheetnames, ["data1", "data2", "empty"])
    ck.expect("multi sheet states", [wb[s].sheet_state for s in wb.sheetnames],
              ["visible", "hidden", "visible"])

    ws1 = wb["data1"]
    ck.expect("data1 dimensions", ws1.dimensions, "A1:B3")
    ck.expect_cell(ws1, "A1", "K", kind=str)
    ck.expect_cell(ws1, "B1", "V", kind=str)
    ck.expect_cell(ws1, "A2", "a", kind=str)
    ck.expect_cell(ws1, "B2", 1, kind=int)
    ck.expect_cell(ws1, "A3", "b", kind=str)
    ck.expect_cell(ws1, "B3", 2, kind=int)

    ws2 = wb["data2"]
    ck.expect("data2 dimensions", ws2.dimensions, "A1:B2")
    ck.expect_cell(ws2, "A1", "X", kind=str)
    ck.expect_cell(ws2, "B1", "Y", kind=str)
    ck.expect_cell(ws2, "A2", "only", kind=str)
    ck.expect_cell(ws2, "B2", False, kind=bool)

    ws3 = wb["empty"]
    ck.expect("empty dimensions", ws3.dimensions, "A1:B1")
    ck.expect_cell(ws3, "A1", "H1", kind=str)
    ck.expect_cell(ws3, "B1", "H2", kind=str)


def check_dimensions(spec: str, ck: Checker) -> None:
    """`path=ROWSxCOLS`: open the workbook read-only and compare sheet extent.

    Only the first sheet is checked: the benchmark outputs this gate targets
    (`spreadsheet-batch.xlsx`, `spreadsheet-streaming.xlsx`) are single-sheet
    workbooks, so `FILE=ROWSxCOLS` describes that sheet.
    """
    try:
        path_str, dims = spec.rsplit("=", 1)
        rows_s, cols_s = dims.lower().split("x")
        rows, cols = int(rows_s), int(cols_s)
    except ValueError:
        ck.failures.append(f"bad --dims spec {spec!r}; expected FILE=ROWSxCOLS")
        return

    import openpyxl

    path = Path(path_str)
    ck.checks += 1
    if not path.exists():
        ck.failures.append(f"{path}: file not found")
        return
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        ck.expect(f"{path.name} first sheet rows", ws.max_row, rows)
        ck.expect(f"{path.name} first sheet cols", ws.max_column, cols)
    finally:
        wb.close()


def main() -> int:
    bootstrap_openpyxl()
    import openpyxl

    dim_specs = sys.argv[1:]
    ck = Checker()

    with tempfile.TemporaryDirectory(prefix="openpyxl-validate-") as tmp:
        out = Path(tmp)
        generate_fixtures(out)

        wb = openpyxl.load_workbook(out / "rs-basic.xlsx")
        validate_basic(wb, ck)
        wb.close()

        wb = openpyxl.load_workbook(out / "rs-multi.xlsx")
        validate_multi(wb, ck)
        wb.close()

        for spec in dim_specs:
            check_dimensions(spec, ck)

    print(f"openpyxl {openpyxl.__version__}: {ck.checks} checks, "
          f"{len(ck.failures)} failures")
    if ck.failures:
        for failure in ck.failures:
            print(f"  FAIL {failure}")
        print("openpyxl validation: FAILED")
        return 1
    print("openpyxl validation: PASS "
          "(parity fixtures value-checked"
          + (f"; {len(dim_specs)} dimension checks" if dim_specs else "") + ")")
    print("note: .xlsb is not readable by openpyxl; it is covered by the "
          "XlsbReader roundtrip tests and the calamine cross-check.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
